from docx import Document
from docx.shared  import Pt  # para aumentar e diminuir o texto
from openpyxl import load_workbook
import os

#caminho do arquivo
name_file = r"caminho onde ta a planilha xlsx"
#abrindo arquivo
sheet_data = load_workbook(name_file)
#selecionando a aba 
sheet_select = sheet_data["Nomes"]


for i in range(2,len(sheet_select["A"])+ 1): 
    # abrindo o arquivo word
    file_word = Document(r"caminho onde ta o docx")
    # tamanho
    style_file = file_word.styles["Normal"]
#nome do aluno que esta na linha 
    name_student = sheet_select['A%s' %i].value
    # subistintuindo o "Nome" e subistituir pelo nome desejado
    for paragraph_file in file_word.paragraphs:
        if "@nome" in paragraph_file.text:
            paragraph_file.text = name_student
            source = style_file.font
            source.name = "Calibri (Corpo)"
            style_file.size = Pt(24)
    path_folder= r"caminho da pasta que deseja salvar\\" + name_student + ".docx"
    # Salva o certificado com o nome do aluno
    file_word.save(path_folder)
print("Certificados gerados com sucesso!")